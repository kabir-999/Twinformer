import hashlib
import re

import numpy as np
import pytest

from app.rag.service import RagService, RagUnavailable, chunk_text
from app.schemas.workloads import RagEvalCase, RagEvalRequest, RagIndexRequest


def bow_factory(model_id, allow_download):
    def embed(texts):
        out = np.zeros((len(texts), 64), dtype="float32")
        for row, text in enumerate(texts):
            for word in re.findall(r"[a-z]+", text.lower()):
                out[row, int(hashlib.md5(word.encode()).hexdigest(), 16) % 64] += 1
        return out / np.maximum(np.linalg.norm(out, axis=1, keepdims=True), 1e-9)
    return embed


@pytest.fixture
def service(tmp_path):
    svc = RagService(str(tmp_path / "docs"), str(tmp_path / "idx"), bow_factory)
    svc.add_document("cats.md", "Cats are small carnivorous mammals. " * 10 + "Whiskers help cats sense nearby objects in darkness. " + "Filler about nothing. " * 30)
    svc.add_document("rockets.txt", "Rockets burn propellant to produce thrust. " * 10 + "Staging discards empty tanks to save mass. " + "Filler about nothing. " * 30)
    return svc


def test_chunk_offsets_point_into_source():
    text = "alpha beta  gamma\ndelta epsilon zeta eta theta"
    chunks = chunk_text("d.md", text, chunk_words=4, overlap_words=1)
    assert all(text[c.char_start:c.char_end] == c.text for c in chunks)
    assert chunks[0].text.split()[-1] == chunks[1].text.split()[0]  # overlap
    assert chunks[-1].text.endswith("theta")


def test_query_before_index_is_unavailable(service):
    with pytest.raises(RagUnavailable):
        service.answer("cats", 3)


def test_index_query_returns_verifiable_citations(service):
    status = service.build_index(RagIndexRequest(chunk_words=40, overlap_words=10))
    assert status.indexed and status.chunk_count > 2 and status.document_count == 2
    answer = service.answer("How do whiskers help cats in darkness?", 3)
    top = answer.citations[0]
    assert top.document == "cats.md"
    source = service.document_text(top.document)
    assert source[top.char_start:top.char_end] == top.text
    assert "word for word" in answer.notes[0]


def test_index_persists_across_restart(service, tmp_path):
    service.build_index(RagIndexRequest(chunk_words=40, overlap_words=10))
    reloaded = RagService(str(tmp_path / "docs"), str(tmp_path / "idx"), bow_factory)
    assert reloaded.status().chunk_count == service.status().chunk_count
    assert reloaded.answer("staging empty tanks", 1).citations[0].document == "rockets.txt"


def test_evaluation_uses_external_labels(service):
    service.build_index(RagIndexRequest(chunk_words=40, overlap_words=10))
    report = service.evaluate(RagEvalRequest(cases=[
        RagEvalCase(question="whiskers sense objects darkness", document="cats.md", must_contain="Whiskers help"),
        RagEvalCase(question="staging empty tanks mass", document="rockets.txt", must_contain="Staging discards"),
        RagEvalCase(question="whiskers sense objects darkness", document="rockets.txt", must_contain="Whiskers help"),  # wrong label -> miss
    ], k_values=[1, 3]))
    assert report.case_count == 3
    assert report.hit_rate_at_k["1"] == pytest.approx(2 / 3, abs=1e-4)
    assert report.per_case[2].first_relevant_rank is None
    assert 0 < report.mrr <= 1


def test_rejects_overlap_not_smaller_than_chunk():
    with pytest.raises(ValueError):
        RagIndexRequest(chunk_words=30, overlap_words=30)


def test_document_name_rejects_path_traversal():
    from app.schemas.workloads import RagDocumentIn
    with pytest.raises(ValueError):
        RagDocumentIn(name="../evil.md", text="x")


def test_folder_sources_with_docx_pptx_and_locations(tmp_path):
    from docx import Document
    from pptx import Presentation
    folder = tmp_path / "Course" / "Ch1"
    folder.mkdir(parents=True)
    doc = Document()
    doc.add_heading("Whiskers", 1)
    doc.add_paragraph("Whiskers help cats sense nearby objects in darkness. " * 5)
    doc.save(folder / "cats.docx")
    deck = Presentation()
    for title in ("Intro slide", "Staging slide"):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "Staging discards empty tanks to save mass. " * 5 if "Staging" in title else "Rockets burn propellant."
    deck.save(folder / "rockets.pptx")
    svc = RagService(str(tmp_path / "docs"), str(tmp_path / "idx"), bow_factory, sources={"Course": tmp_path / "Course"})
    assert svc.list_documents() == ["Course/Ch1/cats.docx", "Course/Ch1/rockets.pptx"]
    status = svc.build_index(RagIndexRequest(chunk_words=40, overlap_words=10))
    assert status.state == "ready" and status.document_count == 2 and not status.documents_changed
    assert status.sources[0].label == "Course" and status.sources[0].file_count == 2
    top = svc.answer("staging discards empty tanks", 1).citations[0]
    assert top.document == "Course/Ch1/rockets.pptx" and "slide 2" in top.location
    assert svc.document_text(top.document)[top.char_start:top.char_end] == top.text
    (folder / "extra.docx").write_bytes((folder / "cats.docx").read_bytes())
    assert svc.status().documents_changed


def test_unreadable_and_empty_files_are_reported_not_fatal(tmp_path):
    folder = tmp_path / "Src"
    folder.mkdir()
    (folder / "broken.pdf").write_bytes(b"not a pdf")
    (folder / "good.txt").write_text("alpha beta gamma delta " * 20)
    svc = RagService(str(tmp_path / "docs"), str(tmp_path / "idx"), bow_factory, sources={"Src": folder})
    svc.build_index(RagIndexRequest(chunk_words=40, overlap_words=10))
    assert svc.status().document_count == 1
    assert any("broken.pdf" in s for s in svc.status().skipped)


def test_bm25_fusion_rescues_exact_keyword_the_embedding_misses(tmp_path):
    """A fake embedder that ignores the rare keyword: dense search alone cannot find it, hybrid can."""
    def blurry_factory(model_id, allow_download):
        def embed(texts):
            out = np.ones((len(texts), 8), dtype="float32")
            for i, t in enumerate(texts):
                out[i, 0] += len(t) % 3  # unrelated to content
            return out / np.linalg.norm(out, axis=1, keepdims=True)
        return embed
    svc = RagService(str(tmp_path / "docs"), str(tmp_path / "idx"), blurry_factory)
    for i in range(30):
        svc.add_document(f"filler{i}.txt", f"General remarks number {i} about nothing in particular. " * 8)
    svc.add_document("target.txt", "The zyzzyva schedule lists exactly one needle fact here. " + "Other words. " * 20)
    svc.build_index(RagIndexRequest(chunk_words=40, overlap_words=10))
    hybrid = svc.retrieve("zyzzyva schedule", 3, hybrid=True)
    assert hybrid[0].document == "target.txt"


def test_retrieval_can_be_limited_to_selected_files(service):
    service.build_index(RagIndexRequest(chunk_words=40, overlap_words=10))
    assert service.indexed_documents() == ["cats.md", "rockets.txt"]
    # the question is about cats, but only the rockets file is allowed
    only_rockets = service.retrieve("whiskers help cats sense objects", 3, documents={"rockets.txt"})
    assert only_rockets and {c.document for c in only_rockets} == {"rockets.txt"}
    both = service.retrieve("whiskers help cats sense objects", 1, documents={"rockets.txt", "cats.md"})
    assert both[0].document == "cats.md"
    with pytest.raises(RagUnavailable):
        service.retrieve("anything", 2, documents={"missing.pdf"})
