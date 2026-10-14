from __future__ import annotations

from pathlib import Path

SUPPORTED = {".pdf", ".pptx", ".docx", ".md", ".txt"}
Segment = tuple[str, str]  # (location label such as "page 3", text)


def _pdf(path: Path) -> list[Segment]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        reader.decrypt("")
    return [(f"page {i + 1}", page.extract_text() or "") for i, page in enumerate(reader.pages)]


def _shape_text(shape) -> list[str]:
    lines: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        lines.append(shape.text_frame.text)
    if getattr(shape, "has_table", False) and shape.has_table:
        lines += [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
    for child in getattr(shape, "shapes", []) if shape.shape_type == 6 else []:  # 6 = group
        lines += _shape_text(child)
    return lines


def _pptx(path: Path) -> list[Segment]:
    from pptx import Presentation
    segments = []
    for i, slide in enumerate(Presentation(str(path)).slides):
        lines = [t for shape in slide.shapes for t in _shape_text(shape)]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        segments.append((f"slide {i + 1}", "\n".join(lines)))
    return segments


def _docx(path: Path) -> list[Segment]:
    from docx import Document
    document = Document(str(path))
    segments: list[Segment] = []
    label, lines = "start", []
    for paragraph in document.paragraphs:
        if paragraph.style is not None and paragraph.style.name.lower().startswith("heading") and paragraph.text.strip():
            if lines:
                segments.append((label, "\n".join(lines)))
            label, lines = paragraph.text.strip()[:60], []
        if paragraph.text.strip():
            lines.append(paragraph.text)
    for table in document.tables:
        lines += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    segments.append((label, "\n".join(lines)))
    return segments


def extract_segments(path: Path) -> list[Segment]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf(path)
    if suffix == ".pptx":
        return _pptx(path)
    if suffix == ".docx":
        return _docx(path)
    return [("text", path.read_text(encoding="utf-8", errors="replace"))]
